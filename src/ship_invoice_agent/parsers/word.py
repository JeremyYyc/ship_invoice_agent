import shutil
import subprocess
from pathlib import Path

from docx import Document

from ship_invoice_agent.exceptions import UnsupportedDocumentError


def parse_docx(path: Path) -> str:
    document = Document(path)
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text.strip())
    for table_index, table in enumerate(document.tables, start=1):
        chunks.append(f"[TABLE {table_index}]")
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(chunks)


def parse_legacy_doc(path: Path) -> str:
    antiword = shutil.which("antiword")
    if not antiword:
        raise UnsupportedDocumentError(
            "legacy .doc parsing requires the 'antiword' executable; convert to .docx or install it"
        )
    result = subprocess.run(
        [antiword, str(path)],
        check=True,
        capture_output=True,
        text=True,
        timeout=60,
    )
    return result.stdout
