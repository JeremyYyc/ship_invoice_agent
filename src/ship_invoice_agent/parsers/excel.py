from pathlib import Path

import xlrd
from openpyxl import load_workbook


def parse_excel(path: Path) -> str:
    if path.suffix.lower() == ".xls":
        return _parse_xls(path)

    workbook = load_workbook(path, data_only=True, read_only=True)
    sections: list[str] = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            continue
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                rows.append(values)
        if not rows:
            continue
        width = max(len(row) for row in rows)
        sections.append(f"[SHEET: {sheet.title}]")
        for row in rows:
            sections.append(" | ".join(row + [""] * (width - len(row))))
    return "\n".join(sections)


def _parse_xls(path: Path) -> str:
    workbook = xlrd.open_workbook(path, on_demand=True)
    sections: list[str] = []
    try:
        for sheet in workbook.sheets():
            if sheet.visibility != 0:
                continue
            rows = [
                ["" if value is None else str(value) for value in sheet.row_values(index)]
                for index in range(sheet.nrows)
            ]
            rows = [row for row in rows if any(row)]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            sections.append(f"[SHEET: {sheet.name}]")
            for row in rows:
                sections.append(" | ".join(row + [""] * (width - len(row))))
    finally:
        workbook.release_resources()
    return "\n".join(sections)
