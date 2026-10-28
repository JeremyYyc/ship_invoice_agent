import hashlib
import json
import re
from collections import Counter
from decimal import Decimal
from pathlib import Path

import pytest
from langchain_core.language_models.fake_chat_models import FakeListChatModel
from openpyxl import load_workbook

from ship_invoice_agent.config import Settings
from ship_invoice_agent.extraction import InvoiceExtractor
from ship_invoice_agent.matching import JsonBillRepository
from ship_invoice_agent.models import Operation
from ship_invoice_agent.output import serialize_operations
from ship_invoice_agent.parsers import DocumentRouter
from ship_invoice_agent.workflow import InvoiceAgent

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
SAMPLES = json.loads((FIXTURES / "office_samples.json").read_text(encoding="utf-8"))


class NoOCRClient:
    def stage_file(self, path):
        raise AssertionError("native Office documents must not be staged for OCR")

    def recognize(self, file_url):
        raise AssertionError("native Office documents must not invoke OCR")


@pytest.mark.parametrize("sample", SAMPLES, ids=lambda sample: sample["name"])
@pytest.mark.parametrize("suffix", ["xlsx", "docx"])
def test_office_router_preserves_invoice_content_without_ocr(sample, suffix):
    path = FIXTURES / "documents" / f"{sample['name']}.{suffix}"
    router = DocumentRouter(Settings(), NoOCRClient())

    documents = router.parse(path)

    assert len(documents) == 1
    raw = documents[0]
    assert raw.source_path == path.resolve()
    assert raw.source_name == path.name
    assert raw.source_type == suffix
    assert raw.file_hash == hashlib.sha256(path.read_bytes()).hexdigest()
    assert raw.metadata == {}
    assert "[SHEET: Charges]" in raw.content if suffix == "xlsx" else "[TABLE 1]" in raw.content
    expected = sample["document"]
    for field in (
        "document_type",
        "reference_no",
        "document_date",
        "issuer",
        "recipient",
        "master_bill_of_lading_no",
        "house_bill_of_lading_no",
        "booking_no",
    ):
        if expected.get(field):
            assert expected[field] in raw.content, f"missing {field} in {path.name}"
    for container in expected.get("container_numbers", []):
        assert container in raw.content
    names = Counter(charge["original_charge_item"] for charge in expected["charges"])
    for name, count in names.items():
        assert raw.content.count(name) >= count
    numeric_values = {
        Decimal(value)
        for value in re.findall(r"(?<![\w.])\d+(?:\.\d+)?(?![\w.])", raw.content.replace(",", ""))
    }
    for charge in expected["charges"]:
        assert Decimal(str(charge["amount"])) in numeric_values
        assert charge["currency"] in raw.content
        if charge.get("payment_term"):
            assert charge["payment_term"] in raw.content
    assert Decimal(str(expected["total"])) in numeric_values


@pytest.mark.parametrize("sample", SAMPLES, ids=lambda sample: sample["name"])
def test_excel_preserves_numeric_amounts_cached_total_and_text_identifier(sample):
    path = FIXTURES / "documents" / f"{sample['name']}.xlsx"
    expected = sample["document"]
    # The production parser uses data_only=True; formula text alone is insufficient.
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook["Charges"]
        for row, charge in enumerate(expected["charges"], start=16):
            assert isinstance(sheet[f"F{row}"].value, (int, float))
            assert sheet[f"F{row}"].value == pytest.approx(charge["amount"])
        total_row = 17 + len(expected["charges"])
        assert sheet[f"F{total_row}"].value == pytest.approx(expected["total"])
        if expected.get("booking_no"):
            assert sheet["B11"].value == "00YYC"
            assert sheet["B11"].data_type == "s"
    finally:
        workbook.close()
    formulas = load_workbook(path, data_only=False, read_only=True)
    try:
        assert formulas["Charges"][f"F{total_row}"].value == (
            f"=SUM(F16:F{15 + len(expected['charges'])})"
        )
    finally:
        formulas.close()


@pytest.mark.parametrize("sample", SAMPLES, ids=lambda sample: sample["name"])
@pytest.mark.parametrize("suffix", ["xlsx", "docx"])
def test_office_workflow_with_fake_model_preserves_charge_rows(
    sample, suffix, tmp_path, wire_response
):
    """Verify native reading and downstream rules; model accuracy requires a separate live run."""
    project = FIXTURES.parents[1]
    path = FIXTURES / "documents" / f"{sample['name']}.{suffix}"
    history_path = tmp_path / "bills.json"
    history_path.write_bytes(b"[]")
    settings = Settings(
        existing_bills_path=history_path,
        charge_aliases_path=project / "data/charge_aliases.json",
        party_aliases_path=project / "data/party_aliases.json",
        currencies_path=project / "data/currencies.json",
        accounting_entities_path=project / "data/accounting_entities.json",
        apply_bill_changes=False,
    )
    raw = DocumentRouter(settings, NoOCRClient()).parse(path)[0]
    model = FakeListChatModel(responses=[json.dumps(wire_response(raw, sample["document"]))])
    agent = InvoiceAgent(
        settings,
        DocumentRouter(settings, NoOCRClient()),
        InvoiceExtractor(model, structured_output_mode="prompt"),
        JsonBillRepository(history_path),
    )

    operations = agent.process(path, accounting_entity=sample["accounting_entity"], apply=False)

    expected = sample["document"]
    assert len(operations) == len(expected["charges"])
    # A previously unknown OCA name now requires review instead of automatic posting.
    expected_operation = (
        Operation.PENDING_REVIEW if sample["name"] == "arrival-notice" else Operation.CREATE
    )
    assert all(operation.operation == expected_operation for operation in operations)
    assert Counter(
        (
            operation.charge_details.charge_item,
            operation.charge_details.amount,
            operation.charge_details.currency,
        )
        for operation in operations
    ) == Counter(
        (charge["charge_item"], Decimal(str(charge["amount"])), charge["currency"])
        for charge in expected["charges"]
    )
    for operation in operations:
        assert operation.document_info.reference_no == expected["reference_no"]
        assert operation.financial_attributes.posting_type == sample["posting_type"]
        assert operation.audit_trail.source_file == path.name
    assert sum(operation.charge_details.amount for operation in operations) == pytest.approx(
        Decimal(str(expected["total"]))
    )
    public = json.loads(serialize_operations(operations))
    assert len(public) == len(expected["charges"])
    assert all(row["operation"] == expected_operation for row in public)
    assert history_path.read_bytes() == b"[]"
